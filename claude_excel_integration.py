"""
Claude Excel Integration
Integração entre Claude AI e planilhas Excel usando a API da Anthropic.
"""

import os
import json
from pathlib import Path
from typing import Any

import anthropic
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ClaudeExcelIntegration:
    """
    Integração entre Claude AI e Excel.

    Permite enviar dados de planilhas para o Claude analisar
    e salvar os resultados de volta em Excel.
    """

    def __init__(self, api_key: str | None = None, model: str = "claude-opus-4-6"):
        self.client = anthropic.Anthropic(api_key=api_key or os.environ["ANTHROPIC_API_KEY"])
        self.model = model

    # ------------------------------------------------------------------
    # Leitura de Excel
    # ------------------------------------------------------------------

    def read_sheet(self, filepath: str, sheet_name: str | None = None) -> list[dict[str, Any]]:
        """
        Lê uma aba do Excel e retorna uma lista de dicts (linha → {header: valor}).
        A primeira linha é tratada como cabeçalho.
        """
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h) if h is not None else f"Col{i+1}" for i, h in enumerate(rows[0])]
        return [
            {headers[i]: cell for i, cell in enumerate(row)}
            for row in rows[1:]
            if any(cell is not None for cell in row)
        ]

    def sheet_to_markdown(self, filepath: str, sheet_name: str | None = None) -> str:
        """Converte uma aba Excel em tabela Markdown para enviar ao Claude."""
        data = self.read_sheet(filepath, sheet_name)
        if not data:
            return "(planilha vazia)"

        headers = list(data[0].keys())
        lines = ["| " + " | ".join(headers) + " |"]
        lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in data:
            lines.append("| " + " | ".join(str(v) if v is not None else "" for v in row.values()) + " |")
        return "\n".join(lines)

    # ------------------------------------------------------------------
    # Chamadas ao Claude
    # ------------------------------------------------------------------

    def analyze_sheet(
        self,
        filepath: str,
        prompt: str,
        sheet_name: str | None = None,
        system: str | None = None,
    ) -> str:
        """
        Envia os dados de uma aba Excel ao Claude junto com um prompt personalizado
        e retorna o texto da resposta.
        """
        table_md = self.sheet_to_markdown(filepath, sheet_name)
        full_prompt = f"{prompt}\n\nDados da planilha:\n\n{table_md}"

        messages: list[anthropic.types.MessageParam] = [
            {"role": "user", "content": full_prompt}
        ]

        kwargs: dict[str, Any] = {
            "model": self.model,
            "max_tokens": 4096,
            "messages": messages,
        }
        if system:
            kwargs["system"] = [
                {
                    "type": "text",
                    "text": system,
                    "cache_control": {"type": "ephemeral"},
                }
            ]

        response = self.client.messages.create(**kwargs)
        return response.content[0].text  # type: ignore[index]

    def analyze_sheet_structured(
        self,
        filepath: str,
        prompt: str,
        json_schema: dict[str, Any],
        sheet_name: str | None = None,
        system: str | None = None,
    ) -> Any:
        """
        Como analyze_sheet, mas pede ao Claude para retornar JSON válido
        conforme o esquema fornecido e faz o parse automaticamente.
        """
        schema_str = json.dumps(json_schema, ensure_ascii=False, indent=2)
        table_md = self.sheet_to_markdown(filepath, sheet_name)
        full_prompt = (
            f"{prompt}\n\n"
            f"Dados da planilha:\n\n{table_md}\n\n"
            f"Retorne SOMENTE um JSON válido seguindo este esquema:\n{schema_str}"
        )

        messages: list[anthropic.types.MessageParam] = [
            {"role": "user", "content": full_prompt}
        ]

        kwargs: dict[str, Any] = {
            "model": self.model,
            "max_tokens": 4096,
            "messages": messages,
        }
        if system:
            kwargs["system"] = [
                {
                    "type": "text",
                    "text": system,
                    "cache_control": {"type": "ephemeral"},
                }
            ]

        response = self.client.messages.create(**kwargs)
        raw = response.content[0].text.strip()  # type: ignore[index]

        # Remove blocos de código markdown se presentes
        if raw.startswith("```"):
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]
        return json.loads(raw)

    def process_rows(
        self,
        filepath: str,
        row_prompt_template: str,
        output_column: str = "claude_response",
        sheet_name: str | None = None,
        system: str | None = None,
    ) -> list[dict[str, Any]]:
        """
        Processa cada linha da planilha individualmente usando o Claude.

        O template de prompt pode referenciar colunas com {nome_da_coluna}.
        Retorna a lista de linhas com a coluna `output_column` adicionada.
        """
        rows = self.read_sheet(filepath, sheet_name)
        results = []

        for row in rows:
            try:
                prompt = row_prompt_template.format(**row)
            except KeyError as exc:
                prompt = f"{row_prompt_template}\n\nDados: {json.dumps(row, ensure_ascii=False)}"
                print(f"Aviso: coluna {exc} não encontrada, usando dados completos.")

            messages: list[anthropic.types.MessageParam] = [
                {"role": "user", "content": prompt}
            ]
            kwargs: dict[str, Any] = {
                "model": self.model,
                "max_tokens": 1024,
                "messages": messages,
            }
            if system:
                kwargs["system"] = system

            response = self.client.messages.create(**kwargs)
            row[output_column] = response.content[0].text.strip()  # type: ignore[index]
            results.append(row)

        return results

    # ------------------------------------------------------------------
    # Escrita em Excel
    # ------------------------------------------------------------------

    def save_results(
        self,
        data: list[dict[str, Any]],
        output_path: str,
        sheet_name: str = "Resultados",
        styled: bool = True,
    ) -> str:
        """
        Salva uma lista de dicts em uma planilha Excel.
        Se `styled=True` aplica formatação básica ao cabeçalho.
        Retorna o caminho do arquivo salvo.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name  # type: ignore[assignment]

        if not data:
            wb.save(output_path)
            return output_path

        headers = list(data[0].keys())

        # Cabeçalho
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            if styled:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="2E75B6")
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Dados
        for row_idx, row in enumerate(data, start=2):
            for col_idx, key in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row.get(key))

        # Auto-ajuste de largura de colunas
        for col_idx, _ in enumerate(headers, start=1):
            max_len = max(
                (
                    len(str(ws.cell(row=r, column=col_idx).value or ""))
                    for r in range(1, len(data) + 2)
                ),
                default=10,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

        wb.save(output_path)
        return output_path

    def append_results_to_sheet(
        self,
        source_filepath: str,
        results: list[dict[str, Any]],
        output_filepath: str | None = None,
        result_column: str = "claude_response",
        source_sheet: str | None = None,
    ) -> str:
        """
        Abre o Excel original e adiciona a coluna de respostas do Claude,
        salvando no arquivo de saída (ou sobrescrevendo o original se None).
        """
        wb = openpyxl.load_workbook(source_filepath)
        ws = wb[source_sheet] if source_sheet else wb.active

        # Descobre o índice da nova coluna
        max_col = ws.max_column
        new_col = max_col + 1
        ws.cell(row=1, column=new_col, value=result_column).font = Font(bold=True)

        for row_idx, result in enumerate(results, start=2):
            ws.cell(row=row_idx, column=new_col, value=result.get(result_column, ""))

        dest = output_filepath or source_filepath
        wb.save(dest)
        return dest

    # ------------------------------------------------------------------
    # Criação de relatório Excel completo a partir de texto do Claude
    # ------------------------------------------------------------------

    def generate_report(
        self,
        prompt: str,
        output_path: str,
        system: str | None = None,
        context_file: str | None = None,
    ) -> str:
        """
        Pede ao Claude para gerar dados tabulares em JSON e salva como Excel.

        O prompt deve solicitar uma resposta no formato:
        {"headers": [...], "rows": [[...], ...], "title": "..."}
        """
        schema = {
            "type": "object",
            "properties": {
                "title": {"type": "string"},
                "headers": {"type": "array", "items": {"type": "string"}},
                "rows": {
                    "type": "array",
                    "items": {"type": "array"},
                },
            },
            "required": ["headers", "rows"],
        }

        context = ""
        if context_file:
            context = f"\n\nContexto da planilha:\n{self.sheet_to_markdown(context_file)}"

        full_prompt = (
            f"{prompt}{context}\n\n"
            f"Retorne SOMENTE um JSON válido com este formato:\n{json.dumps(schema, ensure_ascii=False, indent=2)}"
        )

        messages: list[anthropic.types.MessageParam] = [
            {"role": "user", "content": full_prompt}
        ]
        kwargs: dict[str, Any] = {
            "model": self.model,
            "max_tokens": 4096,
            "messages": messages,
        }
        if system:
            kwargs["system"] = [
                {
                    "type": "text",
                    "text": system,
                    "cache_control": {"type": "ephemeral"},
                }
            ]

        response = self.client.messages.create(**kwargs)
        raw = response.content[0].text.strip()  # type: ignore[index]

        if raw.startswith("```"):
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]

        report_data: dict[str, Any] = json.loads(raw)
        headers: list[str] = report_data["headers"]
        rows: list[list[Any]] = report_data["rows"]
        title: str = report_data.get("title", "Relatório Claude")

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]  # type: ignore[assignment]  # Excel limita a 31 chars

        # Cabeçalho
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="2E75B6")
            cell.alignment = Alignment(horizontal="center")

        # Linhas
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Largura automática
        for col_idx in range(1, len(headers) + 1):
            max_len = max(
                (
                    len(str(ws.cell(row=r, column=col_idx).value or ""))
                    for r in range(1, len(rows) + 2)
                ),
                default=10,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

        wb.save(output_path)
        return output_path
